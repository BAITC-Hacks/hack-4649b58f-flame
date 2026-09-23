from datetime import date
from io import BytesIO

from docx import Document

from app.demo import demo_meeting_result
from app.services.docx_builder import build_docx
from app.ui_state import apply_review_edits


def _document_text(document: Document) -> str:
    paragraphs = [paragraph.text for paragraph in document.paragraphs]
    cells = [cell.text for table in document.tables for row in table.rows for cell in row.cells]
    return "\n".join(paragraphs + cells)


def test_build_docx_opens_and_contains_reviewed_values():
    result = demo_meeting_result("Проверка экспорта", date(2026, 9, 23))
    action_rows = [
        {"Исполнитель": "Ерлан Т.", "Срок": "27.09.2026"},
        {"Исполнитель": "Айгуль", "Срок": "30.09.2026"},
    ]
    transcript_rows = [
        {"Имя говорящего": "Айгуль С."},
        {"Имя говорящего": "Ерлан Т."},
        {"Имя говорящего": "Айгуль С."},
    ]

    reviewed = apply_review_edits(result, action_rows, transcript_rows)
    payload = build_docx(reviewed)
    document = Document(BytesIO(payload))
    text = _document_text(document)

    assert payload.startswith(b"PK")
    assert "Проверка экспорта" in text
    assert "23.09.2026" in text
    assert "Команда согласовала подготовку локального прототипа" in text
    assert "Ерлан Т." in text
    assert "27.09.2026" in text
    assert "Айгуль С." in text
    assert "Подготовить локальную демонстрацию" in text
    assert "Жақсы, жұмаға дейін дайындаймын." in text
    assert reviewed.action_items[0].deadline_iso is None


def test_unchanged_deadline_keeps_normalized_date():
    result = demo_meeting_result("Проверка даты", None)
    original = result.action_items[0].model_copy(
        update={"deadline_text": None, "deadline_iso": date(2026, 9, 27)}
    )
    result = result.model_copy(update={"action_items": [original]})

    reviewed = apply_review_edits(
        result,
        [{"Исполнитель": original.assignee, "Срок": "2026-09-27"}],
        [],
    )

    assert reviewed.action_items[0].deadline_text is None
    assert reviewed.action_items[0].deadline_iso == date(2026, 9, 27)
    document = Document(BytesIO(build_docx(reviewed)))
    assert "27.09.2026" in _document_text(document)
