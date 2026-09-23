from __future__ import annotations

from io import BytesIO

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from app.models.schemas import ActionItem, MeetingResult, TranscriptSegment


_DARK_BLUE = "17365D"
_PALE_BLUE = "EDF3F8"
_BORDER = "D9D9D9"


def _format_time(seconds: float) -> str:
    total_seconds = max(0, int(round(seconds)))
    hours, remainder = divmod(total_seconds, 3600)
    minutes, secs = divmod(remainder, 60)
    if hours:
        return f"{hours:02d}:{minutes:02d}:{secs:02d}"
    return f"{minutes:02d}:{secs:02d}"


def _speaker_label(segment: TranscriptSegment) -> str:
    return segment.speaker_name or segment.speaker_id or "Неизвестный участник"


def _deadline(item: ActionItem) -> str:
    if item.deadline_text:
        return item.deadline_text
    if item.deadline_iso:
        return item.deadline_iso.strftime("%d.%m.%Y")
    return "Не указан"


def _set_cell_shading(cell, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), fill)


def _set_cell_margins(cell, top: int = 100, start: int = 120, bottom: int = 100, end: int = 120) -> None:
    properties = cell._tc.get_or_add_tcPr()
    margins = properties.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        properties.append(margins)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def _set_table_borders(table) -> None:
    properties = table._tbl.tblPr
    borders = properties.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        properties.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "6")
        node.set(qn("w:color"), _BORDER)


def _repeat_header(row) -> None:
    properties = row._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    properties.append(repeat)


def _set_run_font(run, name: str = "Aptos", size: float | None = None) -> None:
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = Pt(size)


def _configure_document(document: Document) -> None:
    section = document.sections[0]
    section.orientation = WD_ORIENT.PORTRAIT
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.7)
    section.right_margin = Inches(0.7)

    normal = document.styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.08

    for style_name, size in (("Title", 24), ("Heading 1", 16), ("Heading 2", 13)):
        style = document.styles[style_name]
        style.font.name = "Aptos Display" if style_name != "Normal" else "Aptos"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.font.bold = style_name != "Title"
        style.paragraph_format.space_before = Pt(12)
        style.paragraph_format.space_after = Pt(6)


def _add_action_items(document: Document, result: MeetingResult) -> None:
    document.add_heading("Поручения", level=1)
    if not result.action_items:
        document.add_paragraph("Поручения не зафиксированы.")
        return

    headers = ("Поручение", "Исполнитель", "Срок", "Уверенность", "Проверка")
    table = document.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    widths = (Inches(2.75), Inches(1.25), Inches(1.15), Inches(0.9), Inches(1.05))
    for index, (cell, header, width) in enumerate(zip(table.rows[0].cells, headers, widths)):
        cell.width = width
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        _set_cell_shading(cell, _DARK_BLUE)
        _set_cell_margins(cell)
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT if index < 3 else WD_ALIGN_PARAGRAPH.CENTER
        run = paragraph.add_run(header)
        run.bold = True
        run.font.color.rgb = RGBColor(255, 255, 255)
        _set_run_font(run, size=9)
    _repeat_header(table.rows[0])

    for row_index, item in enumerate(result.action_items):
        status = "Требует проверки" if item.needs_review else "Без замечаний модели"
        values = (
            item.task,
            item.assignee or "Не указан",
            _deadline(item),
            f"{item.confidence:.0%}",
            status,
        )
        cells = table.add_row().cells
        for column_index, (cell, value, width) in enumerate(zip(cells, values, widths)):
            cell.width = width
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            _set_cell_margins(cell)
            if row_index % 2:
                _set_cell_shading(cell, _PALE_BLUE)
            paragraph = cell.paragraphs[0]
            paragraph.alignment = (
                WD_ALIGN_PARAGRAPH.LEFT if column_index < 3 else WD_ALIGN_PARAGRAPH.CENTER
            )
            run = paragraph.add_run(value)
            _set_run_font(run, size=9)
    _set_table_borders(table)

    for number, item in enumerate(result.action_items, start=1):
        if not item.evidence:
            continue
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.space_before = Pt(5)
        lead = paragraph.add_run(f"Основание {number}: ")
        lead.bold = True
        _set_run_font(lead, size=9.5)
        evidence = paragraph.add_run(item.evidence)
        evidence.italic = True
        _set_run_font(evidence, size=9.5)


def _add_transcript(document: Document, result: MeetingResult) -> None:
    document.add_heading("Транскрипт", level=1)
    if not result.transcript:
        document.add_paragraph("Транскрипт отсутствует.")
        return

    for segment in result.transcript:
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(4)
        stamp = paragraph.add_run(
            f"[{_format_time(segment.start)}-{_format_time(segment.end)}] "
        )
        stamp.font.color.rgb = RGBColor(89, 89, 89)
        _set_run_font(stamp, size=9.5)
        speaker = paragraph.add_run(f"{_speaker_label(segment)}: ")
        speaker.bold = True
        _set_run_font(speaker, size=10)
        text = paragraph.add_run(segment.text)
        _set_run_font(text, size=10)


def build_docx(result: MeetingResult) -> bytes:
    """Build a self-contained meeting protocol DOCX from the shared data model."""

    document = Document()
    _configure_document(document)
    document.core_properties.title = result.title
    document.core_properties.subject = "Протокол совещания"
    document.core_properties.author = "HackAlem AI"

    title = document.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    title_run = title.add_run(result.title)
    title_run.font.color.rgb = RGBColor(0, 0, 0)
    _set_run_font(title_run, name="Aptos Display", size=24)

    date_text = result.meeting_date.strftime("%d.%m.%Y") if result.meeting_date else "Дата не указана"
    metadata = document.add_paragraph()
    metadata.paragraph_format.space_after = Pt(14)
    metadata_run = metadata.add_run(f"Дата совещания: {date_text}")
    metadata_run.font.color.rgb = RGBColor(89, 89, 89)
    _set_run_font(metadata_run, size=10.5)

    document.add_heading("Краткое содержание", level=1)
    document.add_paragraph(result.summary or "Саммари не сформировано.")
    _add_action_items(document, result)

    if result.warnings:
        document.add_heading("Предупреждения", level=1)
        for warning in result.warnings:
            document.add_paragraph(warning, style="List Bullet")

    _add_transcript(document, result)

    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()
